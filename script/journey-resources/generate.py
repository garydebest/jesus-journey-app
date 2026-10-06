"""Generate the Phase 1 client-journey PDFs (and the editable launch kit DOCX)
from the approved manuscripts in this folder.

    python3 script/journey-resources/generate.py

Output: client/public/resources/client-journey/<stage>/<file>
"""
import re
from pathlib import Path

from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Flowable, Image, KeepTogether, ListFlowable, ListItem, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

HERE = Path(__file__).parent
ROOT = HERE.parent.parent
OUT = ROOT / "client/public/resources/client-journey"
LOGO = ROOT / "client/public/assets/logo-mark.png"
FONTS = Path("/usr/share/fonts/truetype/noto")

for name, file in [("Serif", "NotoSerif-Regular.ttf"), ("Serif-SemiBold", "NotoSerif-SemiBold.ttf"),
                   ("Sans", "NotoSans-Regular.ttf"), ("Sans-SemiBold", "NotoSans-SemiBold.ttf"),
                   ("Sans-Bold", "NotoSans-Bold.ttf"), ("Sans-Italic", "NotoSans-Italic.ttf")]:
    pdfmetrics.registerFont(TTFont(name, str(FONTS / file)))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-Bold", italic="Sans-Italic", boldItalic="Sans-Bold")

TEAL = HexColor("#356a65")
INK = HexColor("#1f2a2a")
MUTED = HexColor("#5b6767")
TINT = HexColor("#eef4f3")

S = {
    "eyebrow": ParagraphStyle("eyebrow", fontName="Sans-SemiBold", fontSize=8.5, leading=11, textColor=TEAL, spaceAfter=4),
    "title": ParagraphStyle("title", fontName="Serif-SemiBold", fontSize=22, leading=27, textColor=INK, spaceAfter=14),
    "h2": ParagraphStyle("h2", fontName="Serif-SemiBold", fontSize=13.5, leading=18, textColor=TEAL, spaceBefore=12, spaceAfter=5),
    "h3": ParagraphStyle("h3", fontName="Sans-SemiBold", fontSize=11, leading=15, textColor=INK, spaceBefore=10, spaceAfter=4),
    "body": ParagraphStyle("body", fontName="Sans", fontSize=10, leading=15, textColor=INK, spaceAfter=6, alignment=TA_LEFT),
    "quote": ParagraphStyle("quote", fontName="Sans", fontSize=10, leading=15, textColor=INK),
    "line": ParagraphStyle("line", fontName="Sans", fontSize=10, leading=22, textColor=INK),
}

DOCS = [
    ("coordinator-guide", "prepare/jesus-journey-survey-coordinator-guide.pdf", "Prepare"),
    ("orientation-worksheet", "prepare/jesus-journey-orientation-preparation-worksheet.pdf", "Prepare"),
    ("leadership-briefing", "prepare/jesus-journey-leadership-briefing-guide.pdf", "Prepare"),
    ("readiness-checklist", "prepare/jesus-journey-survey-readiness-checklist.pdf", "Prepare"),
    ("launch-kit", "launch/jesus-journey-launch-communications-kit.pdf", "Launch"),
    ("monitoring-guide", "collect/jesus-journey-response-monitoring-guide.pdf", "Collect"),
    ("closing-checklist", "collect/jesus-journey-survey-closing-checklist.pdf", "Collect"),
    ("debrief-prep", "interpret/jesus-journey-results-debrief-preparation-guide.pdf", "Interpret"),
]


class Box(Flowable):
    """An empty checkbox (font-independent)."""
    def __init__(self, size=9):
        super().__init__(); self.size = size; self.width = size; self.height = size + 3
    def draw(self):
        self.canv.setStrokeColor(TEAL); self.canv.setLineWidth(1)
        self.canv.rect(0, 1, self.size, self.size, stroke=1, fill=0)


CONTENT_W = LETTER[0] - 1.7 * inch


def inline(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    return text.rstrip("  ").rstrip()


def parse(md: str):
    """Tiny parser for the manuscript subset: title, bold headings, ####, lists, checkboxes, quotes."""
    lines = md.split("\n")
    title, blocks, i = None, [], 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        if line.startswith("**Title:**"):
            title = line.replace("**Title:**", "").strip()
        elif line.startswith("#### "):
            blocks.append(("h3", line[5:]))
        elif re.fullmatch(r"\*\*[^*]+\*\*", line.strip()):
            blocks.append(("h2", line.strip()[2:-2]))
        elif line.startswith("> "):
            blocks.append(("quote", line[2:]))
        elif line.startswith("- [ ] "):
            items = []
            while i < len(lines) and lines[i].startswith("- [ ] "):
                items.append(lines[i][6:]); i += 1
            blocks.append(("check", items)); continue
        elif line.startswith("- "):
            items = []
            while i < len(lines) and (lines[i].startswith("- ") or lines[i].startswith("  ")):
                if lines[i].startswith("- "): items.append(lines[i][2:])
                else: items[-1] += "\n" + lines[i].strip()
                i += 1
            blocks.append(("ul", items)); continue
        elif re.match(r"\d+\. ", line):
            items = []
            while i < len(lines) and re.match(r"\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\. ", "", lines[i])); i += 1
            blocks.append(("ol", items)); continue
        else:
            para = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"(- |> |\d+\. |\*\*|#)", lines[i + 1]):
                i += 1; para.append(lines[i].rstrip())
            blocks.append(("p", "<br/>".join(inline(p) for p in para)))
        i += 1
    return title, blocks


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Sans", 8)
    canvas.setFillColor(MUTED)
    canvas.drawString(0.85 * inch, 0.55 * inch, "Jesus Journey · Questions? admin@jesusjourney.life · myjesusjourney.life")
    canvas.drawRightString(LETTER[0] - 0.85 * inch, 0.55 * inch, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf(key: str, rel: str, stage: str):
    title, blocks = parse((HERE / f"{key}.md").read_text())
    story = []
    head = Table([[Image(str(LOGO), width=0.8 * inch, height=0.367 * inch), Paragraph(f"JESUS JOURNEY · {stage.upper()}", S["eyebrow"])]],
                 colWidths=[1.0 * inch, None])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 10), Paragraph(inline(title), S["title"])]
    for kind, val in blocks:
        if kind == "h2":
            story.append(Paragraph(inline(val), S["h2"]))
        elif kind == "h3":
            story.append(Paragraph(inline(val), S["h3"]))
        elif kind == "p":
            story.append(Paragraph(val, S["body"]))
        elif kind == "quote":
            t = Table([[Paragraph(inline(val), S["quote"])]], colWidths=[CONTENT_W])
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), TINT), ("LINEBEFORE", (0, 0), (0, -1), 3, TEAL),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                                   ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
            story += [Spacer(1, 4), t, Spacer(1, 8)]
        elif kind in ("ul", "ol"):
            worksheet = all("___" in it for it in val)
            if worksheet:
                for it in val:
                    story.append(Paragraph(inline(it).replace("\n", "<br/>"), S["line"]))
                continue
            items = [ListItem(Paragraph(inline(it).replace("\n", "<br/>"), S["body"]), leftIndent=14) for it in val]
            story.append(ListFlowable(items, bulletType="1" if kind == "ol" else "bullet", start="1" if kind == "ol" else "•",
                                      bulletFontName="Sans", bulletFontSize=9.5, bulletColor=TEAL, leftIndent=16))
        elif kind == "check":
            rows = [[Box(), Paragraph(inline(it), S["body"])] for it in val]
            t = Table(rows, colWidths=[0.3 * inch, CONTENT_W - 0.3 * inch], hAlign="LEFT")
            t.setStyle(TableStyle([("TOPPADDING", (0, 0), (0, -1), 3), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            story.append(t)
    out = OUT / rel
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out), pagesize=LETTER, leftMargin=0.85 * inch, rightMargin=0.85 * inch,
                            topMargin=0.75 * inch, bottomMargin=0.9 * inch, title=title, author="Jesus Journey",
                            subject=f"Jesus Journey client resource · {stage}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    print("wrote", out.relative_to(ROOT))


def build_docx():
    from docx import Document
    from docx.shared import Pt, RGBColor
    title, blocks = parse((HERE / "launch-kit.md").read_text())
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    h = d.add_heading(title, level=0)
    for r in h.runs: r.font.color.rgb = RGBColor(0x35, 0x6A, 0x65)

    def add_runs(par, text):
        for part in re.split(r"(\*\*.+?\*\*)", text):
            if part.startswith("**") and part.endswith("**"): par.add_run(part[2:-2]).bold = True
            elif part: par.add_run(part)

    for kind, val in blocks:
        raw = re.sub(r"<br/>", "\n", val) if isinstance(val, str) else val
        if isinstance(raw, str):
            raw = raw.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
            raw = re.sub(r"<b>(.+?)</b>", r"**\1**", raw)
        if kind in ("h2", "h3"):
            d.add_heading(raw, level=1 if kind == "h2" else 2)
        elif kind == "p":
            add_runs(d.add_paragraph(), raw)
        elif kind == "quote":
            p = d.add_paragraph(style="Intense Quote"); add_runs(p, raw)
        elif kind in ("ul", "ol", "check"):
            for it in raw:
                add_runs(d.add_paragraph(style="List Number" if kind == "ol" else "List Bullet"), it)
    d.add_paragraph().add_run("Questions? admin@jesusjourney.life · myjesusjourney.life").italic = True
    d.core_properties.title = title; d.core_properties.author = "Jesus Journey"
    out = OUT / "launch/jesus-journey-launch-communications-kit.docx"
    d.save(out)
    print("wrote", out.relative_to(ROOT))


if __name__ == "__main__":
    for key, rel, stage in DOCS:
        build_pdf(key, rel, stage)
    build_docx()
